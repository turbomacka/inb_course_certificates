"""Fyller i platshållare (NAMN, DATUM) i en Word-mall.

Word delar ofta upp en text i flera "runs" (t.ex. "NA" + "MN") och lägger
textrutor både i en modern variant och i en VML-reservvariant. Därför görs
ersättningen per stycke över alla textnoder, i alla XML-delar av dokumentet
(brödtext, textrutor, tabeller, sidhuvud och sidfot).
"""
from docx import Document
from docx.opc.part import XmlPart
from docx.oxml.ns import qn

PLACEHOLDERS = ("NAMN", "DATUM")

W_P = qn("w:p")
W_T = qn("w:t")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _xml_roots(doc):
    """Returnerar rotelementen för alla XML-delar som kan innehålla text."""
    roots = []
    for part in doc.part.package.iter_parts():
        if isinstance(part, XmlPart):
            roots.append(part.element)
    return roots


def _own_texts(paragraph):
    """Textnoder som hör till just detta stycke (inte till nästlade textrutor)."""
    texts = []
    for t in paragraph.iter(W_T):
        parent = t.getparent()
        while parent is not None and parent.tag != W_P:
            parent = parent.getparent()
        if parent is paragraph:
            texts.append(t)
    return texts


def _replace_in_paragraph(paragraph, replacements):
    texts = _own_texts(paragraph)
    if not texts:
        return 0
    count = 0
    for old, new in replacements.items():
        if not old:
            continue
        start = 0
        while True:
            full = "".join(t.text or "" for t in texts)
            idx = full.find(old, start)
            if idx < 0:
                break
            end = idx + len(old)
            pos = 0
            first = True
            for t in texts:
                s = t.text or ""
                t_start, t_end = pos, pos + len(s)
                pos = t_end
                if t_end <= idx or t_start >= end:
                    continue
                a = max(idx, t_start) - t_start
                b = min(end, t_end) - t_start
                # Ersättningen hamnar i den första textnoden och ärver dess formatering.
                t.text = s[:a] + (new if first else "") + s[b:]
                t.set(XML_SPACE, "preserve")
                first = False
            start = idx + len(new)
            count += 1
    return count


def fill_template(template_path, replacements, output_path):
    """Skapar output_path från template_path med ersatta platshållare.

    Returnerar antalet gjorda ersättningar.
    """
    doc = Document(template_path)
    count = 0
    for root in _xml_roots(doc):
        for paragraph in root.iter(W_P):
            count += _replace_in_paragraph(paragraph, replacements)
    doc.save(output_path)
    return count


def find_placeholders(template_path):
    """Returnerar de platshållare (av PLACEHOLDERS) som finns i mallen."""
    doc = Document(template_path)
    found = set()
    for root in _xml_roots(doc):
        for paragraph in root.iter(W_P):
            text = "".join(t.text or "" for t in _own_texts(paragraph))
            for placeholder in PLACEHOLDERS:
                if placeholder in text:
                    found.add(placeholder)
    return [p for p in PLACEHOLDERS if p in found]
