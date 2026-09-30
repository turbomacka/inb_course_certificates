"""Webbversionen (STORAGE_MODE=browser): inget sparas på servern.

Mallar och intyg sparas i användarens webbläsare (IndexedDB, se static/web.js).
Servern fyller bara i mallen och gör PDF:er. Varje förfrågan arbetar i en
tillfällig mapp som tas bort innan svaret skickas.
"""
import io
import json
import os
import tempfile
import zipfile
from datetime import date

from docx.opc.exceptions import PackageNotFoundError
from flask import render_template, request, send_file

from docx_fill import fill_template, find_placeholders
from pdf_convert import ConversionError, convert_to_pdf

# Klienten skickar namnen i omgångar, så att förloppet kan visas och ingen
# enskild förfrågan blir för lång på en liten server.
MAX_NAMES_PER_REQUEST = 10


class BadRequest(Exception):
    pass


def register_browser_routes(app):
    work_root = os.path.join(app.config["DATA_DIR"], "work")
    os.makedirs(work_root, exist_ok=True)

    def page(template):
        return render_template(template, today=date.today().isoformat(), browser_mode=True)

    @app.route("/", methods=["GET"])
    def index():
        return page("web/index.html")

    @app.route("/templates", methods=["GET"])
    def templates():
        return page("web/templates.html")

    @app.route("/certificates", methods=["GET"])
    def certificates():
        return page("web/certificates.html")

    def save_uploaded_template(tmp):
        file = request.files.get("docx_file")
        if not file or not file.filename:
            raise BadRequest("Välj en Word-fil (.docx).")
        if not file.filename.lower().endswith(".docx"):
            raise BadRequest("Endast Word-filer (.docx) kan användas som mall.")
        path = os.path.join(tmp, "mall.docx")
        file.save(path)
        try:
            placeholders = find_placeholders(path)
        except (PackageNotFoundError, KeyError, ValueError):
            raise BadRequest("Filen kunde inte läsas som ett Word-dokument.")
        return path, placeholders

    def run(handler):
        """Kör handler(tmp) i en tillfällig mapp och gör om fel till JSON."""
        try:
            with tempfile.TemporaryDirectory(dir=work_root) as tmp:
                return handler(tmp)
        except BadRequest as exc:
            return {"error": str(exc)}, 400
        except ConversionError as exc:
            return {"error": str(exc)}, 503

    @app.route("/api/inspect", methods=["POST"])
    def api_inspect():
        def handler(tmp):
            _, placeholders = save_uploaded_template(tmp)
            return {"placeholders": placeholders}
        return run(handler)

    @app.route("/api/preview", methods=["POST"])
    def api_preview():
        """PDF av mallen, ifylld med name/datum om de skickas med."""
        def handler(tmp):
            template, _ = save_uploaded_template(tmp)
            source = template
            name = request.form.get("name", "").strip()
            if name:
                source = os.path.join(tmp, "exempel.docx")
                datum = request.form.get("datum", "").strip() or date.today().isoformat()
                fill_template(template, {"NAMN": name, "DATUM": datum}, source)
            pdf = os.path.join(tmp, "forhandsvisning.pdf")
            convert_to_pdf([(source, pdf)], work_root)
            with open(pdf, "rb") as f:
                data = io.BytesIO(f.read())
            return send_file(data, mimetype="application/pdf", download_name="forhandsvisning.pdf", max_age=0)
        return run(handler)

    @app.route("/api/generate", methods=["POST"])
    def api_generate():
        """ZIP med <i>.docx och <i>.pdf för varje namn i ordning (i = 0, 1, ...)."""
        def handler(tmp):
            template, _ = save_uploaded_template(tmp)
            try:
                names = json.loads(request.form.get("names", "[]"))
            except ValueError:
                names = None
            if not isinstance(names, list) or not all(isinstance(n, str) and n.strip() for n in names):
                raise BadRequest("Ogiltig namnlista.")
            if not names:
                raise BadRequest("Ange minst ett namn.")
            if len(names) > MAX_NAMES_PER_REQUEST:
                raise BadRequest(f"Högst {MAX_NAMES_PER_REQUEST} namn per förfrågan.")
            datum = request.form.get("datum", "").strip() or date.today().isoformat()

            jobs = []
            for i, name in enumerate(names):
                docx = os.path.join(tmp, f"{i}.docx")
                fill_template(template, {"NAMN": name.strip(), "DATUM": datum}, docx)
                jobs.append((docx, os.path.join(tmp, f"{i}.pdf")))
            convert_to_pdf(jobs, work_root)

            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for docx, pdf in jobs:
                    zf.write(docx, arcname=os.path.basename(docx))
                    zf.write(pdf, arcname=os.path.basename(pdf))
            buffer.seek(0)
            return send_file(buffer, mimetype="application/zip", download_name="intyg.zip", max_age=0)
        return run(handler)
