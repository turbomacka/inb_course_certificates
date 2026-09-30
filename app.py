import hmac
import io
import logging
import os
import re
import secrets
import shutil
import threading
import zipfile
from datetime import date

from docx.opc.exceptions import PackageNotFoundError
from flask import (Flask, abort, flash, redirect, render_template, request, send_file,
                   session, url_for)
from werkzeug.utils import secure_filename

from docx_fill import fill_template, find_placeholders
from pdf_convert import ConversionError, convert_to_pdf, find_converter
from storage import Storage

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXAMPLE_TEMPLATE = os.path.join(BASE_DIR, "cert_template", "exempelmall.docx")

log = logging.getLogger(__name__)


def _load_secret_key(data_dir):
    """SECRET_KEY från miljön, annars en nyckel som sparas i datakatalogen.

    Nyckeln måste vara densamma i alla Gunicorn-processer, annars tappas
    sessioner och meddelanden mellan förfrågningar.
    """
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    path = os.path.join(data_dir, ".secret_key")
    try:
        with open(path, "x") as f:
            f.write(secrets.token_hex(32))
    except FileExistsError:
        pass
    with open(path) as f:
        return f.read().strip()


def safe_filename_part(value):
    """Gör text säker i ett filnamn men behåller t.ex. å, ä och ö."""
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", value).strip().strip(".")
    return re.sub(r"\s+", " ", value) or "okänd"


def certificate_filename(cert, ext):
    parts = ["Certifikat", cert["kurskod"], cert["student_name"], cert["datum"]]
    return "_".join(safe_filename_part(p) for p in parts if p) + "." + ext


def create_app(config=None):
    app = Flask(__name__)
    app.config.update(
        DATA_DIR=os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data")),
        APP_PASSWORD=os.environ.get("APP_PASSWORD", ""),
        # "server": mallar och intyg sparas i DATA_DIR (desktopversionen).
        # "browser": inget sparas på servern, allt sparas i användarens webbläsare (webbversionen).
        STORAGE_MODE=os.environ.get("STORAGE_MODE", "server"),
        PDF_PREWARM=True,
        MAX_CONTENT_LENGTH=20 * 1024 * 1024,
        # Andra sajter kan inte skicka POST (t.ex. radera intyg) med inloggningen.
        SESSION_COOKIE_SAMESITE="Lax",
    )
    if config:
        app.config.update(config)
    if app.config["STORAGE_MODE"] not in ("server", "browser"):
        raise ValueError(f"Okänt STORAGE_MODE: {app.config['STORAGE_MODE']}")

    os.makedirs(app.config["DATA_DIR"], exist_ok=True)
    app.secret_key = _load_secret_key(app.config["DATA_DIR"])

    # --- gemensamma hjälpfunktioner ---------------------------------------

    def preview_error(message, status):
        return render_template("preview_error.html", message=message), status

    def back_url(default_endpoint):
        """Föregående sida om den ligger på samma webbplats, annars default."""
        ref = request.referrer or ""
        return ref if ref.startswith(request.host_url) else url_for(default_endpoint)

    # --- inloggning (valfri, aktiveras med APP_PASSWORD) ------------------

    @app.before_request
    def require_login():
        if not app.config["APP_PASSWORD"] or session.get("authenticated"):
            return None
        if request.endpoint in ("login", "static"):
            return None
        if request.path.startswith("/api/"):
            return {"error": "Du är inte inloggad."}, 401
        return redirect(url_for("login", next=request.full_path))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if not app.config["APP_PASSWORD"]:
            return redirect(url_for("index"))
        if request.method == "POST":
            # Jämför som bytes: compare_digest godtar bara ASCII i str (å/ä/ö gav 500).
            if hmac.compare_digest(request.form.get("password", "").encode(),
                                   app.config["APP_PASSWORD"].encode()):
                session["authenticated"] = True
                target = request.args.get("next", "")
                # "//x" och "/\x" tolkas av webbläsare som en annan sajt.
                local = target.startswith("/") and not target.startswith(("//", "/\\"))
                return redirect(target if local else url_for("index"))
            flash("Fel lösenord.", "danger")
        return render_template("login.html")

    @app.route("/logout", methods=["POST"])
    def logout():
        session.clear()
        return redirect(url_for("index"))

    @app.route("/download_template")
    def download_template():
        """Exempelmallen (samma adress som i första versionen)."""
        return send_file(EXAMPLE_TEMPLATE, as_attachment=True, download_name="exempelmall.docx")

    @app.errorhandler(413)
    def too_large(_):
        if request.path.startswith("/api/"):
            return {"error": "Filen är för stor (max 20 MB)."}, 413
        flash("Filen är för stor (max 20 MB).", "danger")
        return redirect(back_url("index"))

    if app.config["STORAGE_MODE"] == "browser":
        from web_mode import register_browser_routes
        register_browser_routes(app, EXAMPLE_TEMPLATE)
        return app

    # --- lagring på servern (desktopversionen) ----------------------------

    storage = Storage(app.config["DATA_DIR"])
    app.config["STORAGE"] = storage

    if storage.claim_once("example_template_seeded") and os.path.exists(EXAMPLE_TEMPLATE):
        key = storage.new_key()
        shutil.copyfile(EXAMPLE_TEMPLATE, os.path.join(storage.templates_dir, key + ".docx"))
        storage.add_template("Exempelmall", "exempelmall.docx", key, find_placeholders(EXAMPLE_TEMPLATE))

    def ensure_pdfs(pairs):
        convert_to_pdf(pairs, storage.work_dir)

    def prewarm(pairs):
        """Skapar PDF:er i bakgrunden så att förhandsvisningen går snabbt."""
        if not app.config["PDF_PREWARM"] or not find_converter():
            return

        def run():
            try:
                ensure_pdfs(pairs)
            except Exception:
                log.exception("PDF-förkonvertering misslyckades")

        threading.Thread(target=run, daemon=True).start()

    def pdf_response(pdf_path, download_name, as_attachment=False):
        return send_file(pdf_path, mimetype="application/pdf", as_attachment=as_attachment,
                         download_name=download_name, max_age=0)

    def selected_ids():
        return [int(i) for i in request.form.getlist("ids") if i.isdigit()]

    # --- skapa intyg ------------------------------------------------------

    @app.route("/", methods=["GET"])
    def index():
        templates = storage.list_templates()
        selected = request.args.get("template", type=int)
        if selected is None and templates:
            selected = templates[0]["id"]
        return render_template("index.html", templates=templates, selected=selected,
                               today=date.today().isoformat())

    def _read_generate_form():
        template = storage.get_template(request.form.get("template_id", type=int) or 0)
        kurskod = request.form.get("kurskod", "").strip()
        datum = request.form.get("datum", "").strip() or date.today().isoformat()
        names = [line.strip() for line in request.form.get("student_list", "").splitlines() if line.strip()]
        return template, kurskod, datum, names

    @app.route("/generate", methods=["POST"])
    def generate():
        template, kurskod, datum, names = _read_generate_form()
        if not template:
            flash("Välj en mall.", "danger")
            return redirect(url_for("index"))
        if not names:
            flash("Ange minst ett namn.", "danger")
            return redirect(url_for("index", template=template["id"]))

        students = []
        for name in names:
            key = storage.new_key()
            fill_template(storage.template_docx(template),
                          {"NAMN": name, "DATUM": datum},
                          os.path.join(storage.certificates_dir, key + ".docx"))
            students.append((name, key))
        batch_id, cert_ids = storage.add_batch(kurskod, datum, template["name"], students)

        certs = storage.get_certificates(cert_ids)
        prewarm([(storage.certificate_docx(c), storage.certificate_pdf(c)) for c in certs])
        flash(f"{len(certs)} intyg skapades och har sparats.", "success")
        return redirect(url_for("certificates", batch=batch_id))

    @app.route("/preview-sample", methods=["POST"])
    def preview_sample():
        """Förhandsvisar mallen ifylld med första namnet, utan att spara något."""
        template, _, datum, names = _read_generate_form()
        if not template:
            return preview_error("Välj en mall först.", 400)
        name = names[0] if names else "Förnamn Efternamn"
        key = "sample-" + storage.new_key()
        docx_path = os.path.join(storage.work_dir, key + ".docx")
        pdf_path = os.path.join(storage.work_dir, key + ".pdf")
        try:
            fill_template(storage.template_docx(template),
                          {"NAMN": name, "DATUM": datum}, docx_path)
            ensure_pdfs([(docx_path, pdf_path)])
            with open(pdf_path, "rb") as f:
                data = io.BytesIO(f.read())
        except ConversionError as exc:
            return preview_error(str(exc), 503)
        finally:
            for path in (docx_path, pdf_path):
                if os.path.exists(path):
                    os.remove(path)
        return send_file(data, mimetype="application/pdf", download_name="forhandsvisning.pdf", max_age=0)

    # --- mallar -----------------------------------------------------------

    @app.route("/templates", methods=["GET"])
    def templates():
        return render_template("templates.html", templates=storage.list_templates())

    @app.route("/templates", methods=["POST"])
    def upload_template():
        back = request.form.get("next") == "index"
        file = request.files.get("docx_file")
        if not file or not file.filename:
            flash("Välj en Word-fil (.docx) att ladda upp.", "danger")
            return redirect(url_for("index" if back else "templates"))
        if not file.filename.lower().endswith(".docx"):
            flash("Endast Word-filer (.docx) kan användas som mall.", "danger")
            return redirect(url_for("index" if back else "templates"))

        key = storage.new_key()
        path = os.path.join(storage.templates_dir, key + ".docx")
        file.save(path)
        try:
            placeholders = find_placeholders(path)
        except (PackageNotFoundError, KeyError, ValueError):
            os.remove(path)
            flash("Filen kunde inte läsas som ett Word-dokument.", "danger")
            return redirect(url_for("index" if back else "templates"))

        name = request.form.get("name", "").strip() or os.path.splitext(file.filename)[0]
        template_id = storage.add_template(name, secure_filename(file.filename) or "mall.docx", key, placeholders)
        template = storage.get_template(template_id)
        prewarm([(storage.template_docx(template), storage.template_pdf(template))])

        if "NAMN" not in placeholders:
            flash(f"Mallen ”{name}” sparades, men den innehåller ingen platshållare NAMN.", "warning")
        else:
            flash(f"Mallen ”{name}” sparades.", "success")
        return redirect(url_for("index", template=template_id) if back else url_for("templates"))

    def _get_template_or_404(template_id):
        template = storage.get_template(template_id)
        if not template:
            abort(404)
        return template

    @app.route("/templates/<int:template_id>/preview")
    def template_preview(template_id):
        template = _get_template_or_404(template_id)
        pdf_path = storage.template_pdf(template)
        try:
            ensure_pdfs([(storage.template_docx(template), pdf_path)])
        except ConversionError as exc:
            return preview_error(str(exc), 503)
        return pdf_response(pdf_path, safe_filename_part(template["name"]) + ".pdf")

    @app.route("/templates/<int:template_id>/download")
    def template_download(template_id):
        template = _get_template_or_404(template_id)
        return send_file(storage.template_docx(template), as_attachment=True,
                         download_name=safe_filename_part(template["name"]) + ".docx")

    @app.route("/templates/<int:template_id>/delete", methods=["POST"])
    def template_delete(template_id):
        template = _get_template_or_404(template_id)
        storage.delete_template(template_id)
        flash(f"Mallen ”{template['name']}” togs bort. Redan skapade intyg finns kvar.", "success")
        return redirect(url_for("templates"))

    # --- sparade intyg ----------------------------------------------------

    @app.route("/certificates", methods=["GET"])
    def certificates():
        batch_id = request.args.get("batch", type=int)
        query = request.args.get("q", "").strip()
        return render_template("certificates.html",
                               certificates=storage.list_certificates(batch_id, query),
                               batches=storage.list_batches(), batch_id=batch_id, query=query)

    def _get_certificate_or_404(cert_id):
        cert = storage.get_certificate(cert_id)
        if not cert:
            abort(404)
        return cert

    @app.route("/certificates/<int:cert_id>/preview")
    def certificate_preview(cert_id):
        cert = _get_certificate_or_404(cert_id)
        try:
            ensure_pdfs([(storage.certificate_docx(cert), storage.certificate_pdf(cert))])
        except ConversionError as exc:
            return preview_error(str(exc), 503)
        return pdf_response(storage.certificate_pdf(cert), certificate_filename(cert, "pdf"))

    @app.route("/certificates/<int:cert_id>/download/<fmt>")
    def certificate_download(cert_id, fmt):
        cert = _get_certificate_or_404(cert_id)
        if fmt == "docx":
            return send_file(storage.certificate_docx(cert), as_attachment=True,
                             download_name=certificate_filename(cert, "docx"))
        if fmt != "pdf":
            abort(404)
        try:
            ensure_pdfs([(storage.certificate_docx(cert), storage.certificate_pdf(cert))])
        except ConversionError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("certificates", batch=cert["batch_id"]))
        return pdf_response(storage.certificate_pdf(cert), certificate_filename(cert, "pdf"), as_attachment=True)

    @app.route("/certificates/download", methods=["POST"])
    def certificates_download():
        fmt = request.form.get("fmt", "pdf")
        certs = storage.get_certificates(selected_ids())
        if fmt not in ("pdf", "docx") or not certs:
            flash("Markera minst ett intyg att ladda ner.", "warning")
            return redirect(back_url("certificates"))
        if len(certs) == 1:
            return redirect(url_for("certificate_download", cert_id=certs[0]["id"], fmt=fmt))

        if fmt == "pdf":
            try:
                ensure_pdfs([(storage.certificate_docx(c), storage.certificate_pdf(c)) for c in certs])
            except ConversionError as exc:
                flash(str(exc), "danger")
                return redirect(back_url("certificates"))

        buffer = io.BytesIO()
        used = set()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for cert in certs:
                name = certificate_filename(cert, fmt)
                stem, n = name[: -len(fmt) - 1], 2
                while name in used:
                    name, n = f"{stem} ({n}).{fmt}", n + 1
                used.add(name)
                path = storage.certificate_pdf(cert) if fmt == "pdf" else storage.certificate_docx(cert)
                zf.write(path, arcname=name)
        buffer.seek(0)
        return send_file(buffer, mimetype="application/zip", as_attachment=True,
                         download_name=f"intyg_{fmt}_{date.today().isoformat()}.zip")

    @app.route("/certificates/delete", methods=["POST"])
    def certificates_delete():
        return _delete_and_redirect(selected_ids())

    @app.route("/certificates/<int:cert_id>/delete", methods=["POST"])
    def certificate_delete(cert_id):
        return _delete_and_redirect([cert_id])

    def _delete_and_redirect(ids):
        deleted = storage.delete_certificates(ids)
        if deleted:
            flash(f"{deleted} intyg togs bort.", "success")
        else:
            flash("Markera minst ett intyg att ta bort.", "warning")
        return redirect(back_url("certificates"))

    return app


app = create_app()

if __name__ == "__main__":
    app.run(debug=True)
