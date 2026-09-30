import zipfile

from docx import Document

from conftest import TEXTBOX_TEMPLATE
from docx_fill import fill_template, find_placeholders

REPL = {"NAMN": "Åsa Öberg", "DATUM": "2026-09-30"}


def all_xml_text(path):
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xml"))


def test_replaces_placeholder_split_over_runs_keeps_first_run_format(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    p = doc.add_paragraph("Intyg för ")
    bold = p.add_run("NA")
    bold.bold = True
    p.add_run("MN, datum DA")
    p.add_run("TUM")
    doc.save(src)

    assert fill_template(src, REPL, out) == 2
    para = Document(out).paragraphs[0]
    assert para.text == "Intyg för Åsa Öberg, datum 2026-09-30"
    assert para.runs[1].text == "Åsa Öberg" and para.runs[1].bold


def test_replaces_in_tables_headers_and_footers(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Datum: DATUM"
    doc.sections[0].header.paragraphs[0].text = "NAMN"
    doc.sections[0].footer.paragraphs[0].text = "DATUM"
    doc.save(src)

    fill_template(src, REPL, out)
    result = Document(out)
    assert result.tables[0].cell(0, 0).text == "Datum: 2026-09-30"
    assert result.sections[0].header.paragraphs[0].text == "Åsa Öberg"
    assert result.sections[0].footer.paragraphs[0].text == "2026-09-30"


def test_replacement_containing_placeholder_does_not_loop(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("NAMN och NAMN")
    doc.save(src)
    fill_template(src, {"NAMN": "NAMNsson"}, out)
    assert Document(out).paragraphs[0].text == "NAMNsson och NAMNsson"


def test_textboxes_including_vml_fallback_are_filled(tmp_path):
    out = tmp_path / "out.docx"
    assert "NAMN" in all_xml_text(TEXTBOX_TEMPLATE)
    fill_template(TEXTBOX_TEMPLATE, REPL, out)
    text = all_xml_text(out)
    assert "NAMN" not in text and "DATUM" not in text
    assert text.count("Åsa Öberg") >= 2  # både textruta och VML-reserv


def test_find_placeholders():
    assert find_placeholders(TEXTBOX_TEMPLATE) == ["NAMN", "DATUM"]
