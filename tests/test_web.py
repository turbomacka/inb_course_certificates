import io
import json
import os
import zipfile

import pytest
from docx import Document

import pdf_convert
from app import create_app
from conftest import EXAMPLE

needs_converter = pytest.mark.skipif(not pdf_convert.find_converter(), reason="Varken Word eller LibreOffice finns")


@pytest.fixture
def data_dir(tmp_path):
    return tmp_path / "data"


@pytest.fixture
def client(data_dir):
    return create_app({"DATA_DIR": str(data_dir), "STORAGE_MODE": "browser", "TESTING": True}).test_client()


def template_upload(path=EXAMPLE, filename="mall.docx"):
    with open(path, "rb") as f:
        return (io.BytesIO(f.read()), filename)


def stored_files(data_dir):
    """Alla filer under DATA_DIR utom sessionsnyckeln, låsfilen och LibreOffice-profilen."""
    found = []
    for root, _, files in os.walk(data_dir):
        if "lo-profile" in root:
            continue
        found += [f for f in files if f not in (".secret_key", "soffice.lock")]
    return found


def test_pages_render_without_server_storage(client, data_dir):
    for url in ("/", "/templates", "/certificates", "/download_template"):
        assert client.get(url).status_code == 200
    assert "web.js" in client.get("/certificates").get_data(as_text=True)
    assert not (data_dir / "app.db").exists()


def test_server_storage_routes_are_absent(client):
    assert client.get("/certificates/1/preview").status_code == 404
    assert client.post("/generate", data={}).status_code in (404, 405)


def test_inspect_returns_placeholders(client, data_dir):
    response = client.post("/api/inspect", data={"docx_file": template_upload()}, content_type="multipart/form-data")
    assert response.get_json() == {"placeholders": ["NAMN", "DATUM"]}
    assert stored_files(data_dir) == []


def test_inspect_rejects_bad_files(client):
    for upload in ((io.BytesIO(b"x"), "a.pdf"), (io.BytesIO(b"not a zip"), "a.docx")):
        response = client.post("/api/inspect", data={"docx_file": upload}, content_type="multipart/form-data")
        assert response.status_code == 400 and response.get_json()["error"]
    assert client.post("/api/inspect", data={}).status_code == 400


def test_generate_validates_names(client):
    for names in ("", "inte json", json.dumps(["ok", ""]), json.dumps([]), json.dumps(["x"] * 11)):
        response = client.post("/api/generate", data={"docx_file": template_upload(), "names": names},
                               content_type="multipart/form-data")
        assert response.status_code == 400, names


def test_generate_without_converter_gives_503_and_leaves_nothing(client, data_dir, monkeypatch):
    monkeypatch.setattr(pdf_convert, "find_converter", lambda: None)
    response = client.post("/api/generate", data={"docx_file": template_upload(), "names": json.dumps(["Anna"])},
                           content_type="multipart/form-data")
    assert response.status_code == 503 and "Word eller LibreOffice" in response.get_json()["error"]
    assert stored_files(data_dir) == []


@needs_converter
def test_generate_returns_docx_and_pdf_and_stores_nothing(client, data_dir, tmp_path):
    names = ["Åsa Öberg", "Bo Berg"]
    response = client.post("/api/generate", data={"docx_file": template_upload(), "datum": "2026-09-30",
                                                  "names": json.dumps(names)}, content_type="multipart/form-data")
    assert response.status_code == 200 and response.mimetype == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(response.data))
    assert sorted(zf.namelist()) == ["0.docx", "0.pdf", "1.docx", "1.pdf"]
    assert zf.read("0.pdf").startswith(b"%PDF")
    out = tmp_path / "0.docx"
    out.write_bytes(zf.read("0.docx"))
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Åsa Öberg" in text and "2026-09-30" in text
    assert stored_files(data_dir) == []


@needs_converter
def test_preview_template_and_sample(client, data_dir):
    raw = client.post("/api/preview", data={"docx_file": template_upload()}, content_type="multipart/form-data")
    sample = client.post("/api/preview", data={"docx_file": template_upload(), "name": "Anna", "datum": "2026-09-30"},
                         content_type="multipart/form-data")
    for response in (raw, sample):
        assert response.status_code == 200 and response.data.startswith(b"%PDF")
    assert stored_files(data_dir) == []


def test_api_requires_login_with_json_401(data_dir):
    app = create_app({"DATA_DIR": str(data_dir), "STORAGE_MODE": "browser", "APP_PASSWORD": "hemligt"})
    client = app.test_client()
    response = client.post("/api/inspect", data={"docx_file": template_upload()}, content_type="multipart/form-data")
    assert response.status_code == 401 and response.get_json()["error"]
    assert client.get("/").headers["Location"].startswith("/login")
    client.post("/login", data={"password": "hemligt"})
    response = client.post("/api/inspect", data={"docx_file": template_upload()}, content_type="multipart/form-data")
    assert response.status_code == 200


def test_unknown_storage_mode_is_rejected(data_dir):
    with pytest.raises(ValueError):
        create_app({"DATA_DIR": str(data_dir), "STORAGE_MODE": "moln"})
