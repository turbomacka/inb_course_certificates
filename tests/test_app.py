import io
import os
import zipfile

import pytest
from docx import Document

import pdf_convert
from app import create_app, safe_filename_part
from conftest import EXAMPLE

needs_converter = pytest.mark.skipif(not pdf_convert.find_converter(), reason="Varken Word eller LibreOffice finns")


def upload(client, path=EXAMPLE, name="Testmall", filename="mall.docx", **extra):
    with open(path, "rb") as f:
        data = {"name": name, "docx_file": (io.BytesIO(f.read()), filename), **extra}
    return client.post("/templates", data=data, content_type="multipart/form-data")


def generate(client, template_id, names="Anna Andersson\nBertil Berg", kurskod="HPC012"):
    return client.post("/generate", data={"template_id": template_id, "kurskod": kurskod,
                                          "datum": "2026-09-30", "student_list": names})


def test_example_template_is_seeded_once(tmp_path):
    config = {"DATA_DIR": str(tmp_path / "d"), "PDF_PREWARM": False}
    storage = create_app(config).config["STORAGE"]
    assert [t["name"] for t in storage.list_templates()] == ["Exempelmall"]
    storage.delete_template(storage.list_templates()[0]["id"])
    assert create_app(config).config["STORAGE"].list_templates() == []


def test_pages_render(client):
    for url in ("/", "/templates", "/certificates", "/download_template"):
        assert client.get(url).status_code == 200


def test_upload_template_is_saved_and_listed(client, storage):
    response = upload(client)
    assert response.status_code == 302
    template = next(t for t in storage.list_templates() if t["name"] == "Testmall")
    assert template["placeholders"] == "NAMN,DATUM"
    assert os.path.exists(storage.template_docx(template))
    assert "Testmall" in client.get("/templates").get_data(as_text=True)


def test_upload_from_index_redirects_back_with_template_selected(client, storage):
    response = upload(client, next="index")
    template = next(t for t in storage.list_templates() if t["name"] == "Testmall")
    assert response.headers["Location"].endswith(f"/?template={template['id']}")


def test_upload_rejects_non_docx(client, storage):
    before = len(storage.list_templates())
    client.post("/templates", data={"docx_file": (io.BytesIO(b"x"), "a.pdf")}, content_type="multipart/form-data")
    client.post("/templates", data={"docx_file": (io.BytesIO(b"not a zip"), "a.docx")},
                content_type="multipart/form-data")
    assert len(storage.list_templates()) == before
    assert os.listdir(storage.templates_dir) == [f"{storage.list_templates()[0]['file_key']}.docx"]


def test_generate_saves_certificates(client, storage):
    template_id = storage.list_templates()[0]["id"]
    response = generate(client, template_id)
    assert response.status_code == 302 and "/certificates?batch=" in response.headers["Location"]
    certs = storage.list_certificates()
    assert [c["student_name"] for c in certs] == ["Anna Andersson", "Bertil Berg"]
    body = "\n".join(p.text for p in Document(storage.certificate_docx(certs[0])).paragraphs)
    assert "Anna Andersson" in body and "2026-09-30" in body
    page = client.get(response.headers["Location"]).get_data(as_text=True)
    assert "Anna Andersson" in page and "Bertil Berg" in page


def test_generate_requires_names(client, storage):
    generate(client, storage.list_templates()[0]["id"], names="  \n ")
    assert storage.list_certificates() == []


def test_kurskod_is_only_used_in_filename(client, storage, tmp_path):
    src = tmp_path / "kurskod.docx"
    doc = Document()
    doc.add_paragraph("NAMN KURSKOD")
    doc.save(src)
    upload(client, path=src, name="Med kurskod")
    template = next(t for t in storage.list_templates() if t["name"] == "Med kurskod")
    assert template["placeholders"] == "NAMN"

    generate(client, template["id"], names="Anna")
    cert = storage.list_certificates()[0]
    assert Document(storage.certificate_docx(cert)).paragraphs[0].text == "Anna KURSKOD"
    assert "Certifikat_HPC012_Anna_2026-09-30.docx" in \
        client.get(f"/certificates/{cert['id']}/download/docx").headers["Content-Disposition"]


def test_download_single_docx_with_readable_filename(client, storage):
    generate(client, storage.list_templates()[0]["id"], names="Åsa Öberg")
    cert = storage.list_certificates()[0]
    response = client.get(f"/certificates/{cert['id']}/download/docx")
    assert response.status_code == 200
    assert "Certifikat_HPC012_%C3%85sa%20%C3%96berg_2026-09-30.docx" in response.headers["Content-Disposition"]


def test_download_selected_as_zip_with_unique_names(client, storage):
    generate(client, storage.list_templates()[0]["id"], names="Anna\nAnna\nBo")
    ids = [c["id"] for c in storage.list_certificates()]
    response = client.post("/certificates/download", data={"ids": ids[:2], "fmt": "docx"})
    assert response.mimetype == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(response.data)).namelist()
    assert names == ["Certifikat_HPC012_Anna_2026-09-30.docx", "Certifikat_HPC012_Anna_2026-09-30 (2).docx"]


def test_download_single_selected_redirects_to_file(client, storage):
    generate(client, storage.list_templates()[0]["id"], names="Anna")
    cert_id = storage.list_certificates()[0]["id"]
    response = client.post("/certificates/download", data={"ids": [cert_id], "fmt": "docx"})
    assert response.headers["Location"].endswith(f"/certificates/{cert_id}/download/docx")


def test_delete_single_and_selected(client, storage):
    generate(client, storage.list_templates()[0]["id"], names="A\nB\nC")
    certs = storage.list_certificates()
    client.post(f"/certificates/{certs[0]['id']}/delete")
    assert not os.path.exists(storage.certificate_docx(certs[0]))
    client.post("/certificates/delete", data={"ids": [certs[1]["id"], certs[2]["id"]]})
    assert storage.list_certificates() == []
    assert storage.list_batches() == []
    assert os.listdir(storage.certificates_dir) == []


def test_filter_by_batch_and_search(client, storage):
    template_id = storage.list_templates()[0]["id"]
    generate(client, template_id, names="Anna", kurskod="AAA")
    generate(client, template_id, names="Bo", kurskod="BBB")
    batches = storage.list_batches()
    assert [c["student_name"] for c in storage.list_certificates(batches[0]["id"])] == ["Bo"]
    assert [c["student_name"] for c in storage.list_certificates(query="ann")] == ["Anna"]


def test_deleting_template_keeps_certificates(client, storage):
    template = storage.list_templates()[0]
    generate(client, template["id"], names="Anna")
    client.post(f"/templates/{template['id']}/delete")
    assert storage.list_templates() == []
    cert = storage.list_certificates()[0]
    assert client.get(f"/certificates/{cert['id']}/download/docx").status_code == 200


def test_preview_without_converter_shows_message(client, storage, monkeypatch):
    monkeypatch.setattr(pdf_convert, "find_converter", lambda: None)
    template_id = storage.list_templates()[0]["id"]
    response = client.get(f"/templates/{template_id}/preview")
    assert response.status_code == 503 and "Word eller LibreOffice" in response.get_data(as_text=True)
    response = client.post("/preview-sample", data={"template_id": template_id, "student_list": "Anna"})
    assert response.status_code == 503
    assert [f for f in os.listdir(storage.work_dir) if f.startswith("sample-")] == []


def test_unknown_ids_give_404(client):
    assert client.get("/certificates/999/preview").status_code == 404
    assert client.get("/templates/999/download").status_code == 404
    assert client.get("/certificates/1/download/exe").status_code == 404


def test_password_protection(tmp_path):
    app = create_app({"DATA_DIR": str(tmp_path / "d"), "APP_PASSWORD": "hemligt", "PDF_PREWARM": False})
    client = app.test_client()
    assert client.get("/certificates").headers["Location"].startswith("/login")
    client.post("/login", data={"password": "fel"})
    assert client.get("/certificates").status_code == 302
    response = client.post("/login?next=/certificates", data={"password": "hemligt"})
    assert response.headers["Location"] == "/certificates"
    assert client.get("/certificates").status_code == 200


def test_login_handles_non_ascii_and_rejects_external_next(tmp_path):
    app = create_app({"DATA_DIR": str(tmp_path / "d"), "APP_PASSWORD": "blåbär", "PDF_PREWARM": False})
    client = app.test_client()
    assert client.post("/login", data={"password": "fel å"}).status_code == 200
    for target in ("//evil.example", "/\\evil.example", "https://evil.example"):
        response = client.post(f"/login?next={target}", data={"password": "blåbär"})
        assert response.headers["Location"] == "/"
    assert "SameSite=Lax" in response.headers["Set-Cookie"]


def test_safe_filename_part():
    assert safe_filename_part('Åsa / "Ö" : Berg ') == "Åsa Ö Berg"
    assert safe_filename_part("../..") == "okänd"


@needs_converter
def test_previews_and_pdf_zip(client, storage):
    template_id = storage.list_templates()[0]["id"]
    response = client.get(f"/templates/{template_id}/preview")
    assert response.status_code == 200 and response.data.startswith(b"%PDF")

    response = client.post("/preview-sample", data={"template_id": template_id, "kurskod": "X",
                                                    "datum": "2026-09-30", "student_list": "Anna"})
    assert response.data.startswith(b"%PDF")

    generate(client, template_id, names="\n".join(f"Person {i}" for i in range(7)))
    certs = storage.list_certificates()
    assert client.get(f"/certificates/{certs[0]['id']}/preview").data.startswith(b"%PDF")
    response = client.post("/certificates/download", data={"ids": [c["id"] for c in certs], "fmt": "pdf"})
    zf = zipfile.ZipFile(io.BytesIO(response.data))
    assert len(zf.namelist()) == 7
    assert all(zf.read(n).startswith(b"%PDF") for n in zf.namelist())
